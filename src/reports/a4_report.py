from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "17324D"
BLUE = "277DA1"
PALE = "EAF2F8"


def populate_a4_sheet(ws: Worksheet, report: dict[str, Any]) -> None:
    tournament = report["tournament"]
    summary = report["summary"]
    ws.title = "A4 overzicht"
    ws.merge_cells("A1:H1")
    ws["A1"] = tournament["name"]
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:H2")
    ws["A2"] = f"{tournament['location']} | {tournament['start_date']} | {summary['teams']} teams"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(color="52606D")

    cards = [("Spelers", summary["players"]), ("Gem. leeftijd", summary["avg_age"]), ("Gem. lengte", summary["avg_height"]), ("Teams", summary["teams"])]
    for index, (label, value) in enumerate(cards):
        col = 1 + index * 2
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 1)
        label_cell = ws.cell(4, col, label)
        value_cell = ws.cell(5, col, value if value is not None else "n.a.")
        label_cell.fill = PatternFill("solid", fgColor=BLUE)
        label_cell.font = Font(bold=True, color="FFFFFF")
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.fill = PatternFill("solid", fgColor=PALE)
        value_cell.font = Font(size=16, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")

    headers = ["Land", "Ranking", "# spelers", "Gem. leeftijd", "Gem. lengte netspelers"]
    for col, header in enumerate(headers, 1):
        ws.cell(8, col, header)
    result_map = {row["country_code"]: row for row in report["results"]}
    for row_index, country in enumerate(report["countries"], 9):
        result = result_map[country["country_code"]]
        values = [country["country_code"], result["ranking"], country["players"], country["avg_age"], country["avg_net_height"]]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)

    ws["G8"] = "Geboortekwartaal"
    ws["H8"] = "#"
    for row_index, quarter in enumerate(("Q1", "Q2", "Q3", "Q4"), 9):
        ws.cell(row_index, 7, quarter)
        ws.cell(row_index, 8, report["quarters"][quarter])

    ws["G14"] = "Ranking-analyse"
    ws["H14"] = "Waarde"
    ws["G15"] = "Ranking-lengte"
    ws["H15"] = report["correlations"]["ranking_height"]
    ws["G16"] = "Ranking-leeftijd"
    ws["H16"] = report["correlations"]["ranking_age"]
    ws["H15"].number_format = "0.000"
    ws["H16"].number_format = "0.000"
    ws.merge_cells("A19:H21")
    ws["A19"] = "Een negatieve correlatie betekent dat een hogere gemiddelde waarde samenhangt met een betere (lagere) ranking. Correlatie beschrijft samenhang, geen oorzaak."
    ws["A19"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A19"].font = Font(italic=True, color="52606D")
    ws.row_dimensions[19].height = 24
    for header_row in (8, 14):
        for cell in ws[header_row]:
            if cell.value is not None:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.font = Font(bold=True, color="FFFFFF")
    for column, width in {"A": 13, "B": 12, "C": 13, "D": 16, "E": 24, "F": 3, "G": 21, "H": 13}.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:H21"
