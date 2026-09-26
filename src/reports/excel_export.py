from __future__ import annotations

from io import BytesIO
from math import ceil, floor
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17324D"
BLUE = "277DA1"
LIGHT = "EAF2F8"
ORANGE = "F4A261"
WHITE = "FFFFFF"
PALETTE = ["277DA1", "F94144", "90BE6D", "8E6BBE", "F8961E", "43AA8B"]


def create_excel_report(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overzicht"
    players = workbook.create_sheet("Spelers")
    results = workbook.create_sheet("Resultaten")
    _overview(overview, report)
    _players(players, report)
    _results(results, report)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _overview(ws, report: dict[str, Any]) -> None:
    tournament = report["tournament"]
    summary = report["summary"]
    ws.merge_cells("A1:N2")
    ws["A1"] = tournament["name"]
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    details = [
        ("Toernooi", tournament["name"]), ("Categorie", tournament["category"]),
        ("Geslacht", tournament["gender"]), ("Locatie", tournament["location"]),
        ("Startdatum", tournament["start_date"]), ("Bron", tournament.get("source_file") or ""),
    ]
    for row, (label, value) in enumerate(details, 4):
        ws.cell(row, 1, label).font = Font(bold=True, color=NAVY)
        ws.cell(row, 2, value)
    cards = [("Teams", summary["teams"]), ("Spelers", summary["players"]), ("Gem. leeftijd", summary["avg_age"]), ("Gem. lengte", summary["avg_height"])]
    for idx, (label, value) in enumerate(cards):
        col = 4 + idx * 3
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=7, end_column=col + 1)
        ws.cell(4, col, label)
        ws.cell(5, col, value if value is not None else "n.a.")
        for cell in (ws.cell(4, col), ws.cell(5, col)):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(4, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(4, col).font = Font(bold=True, color=WHITE)
        ws.cell(5, col).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(5, col).font = Font(size=16, bold=True, color=NAVY)

    result_map = {row["country_code"]: row for row in report["results"]}
    country_start = 10
    country_headers = ["Land", "Ranking", "# spelers", "Gem. leeftijd", "Gem. lengte netspelers", "Q1", "Q2", "Q3", "Q4"]
    country_rows = []
    for row in report["countries"]:
        country_rows.append([
            row["country_code"], result_map.get(row["country_code"], {}).get("ranking"), row["players"],
            row["avg_age"], row["avg_net_height"], row["Q1"], row["Q2"], row["Q3"], row["Q4"],
        ])
    _write_table(ws, country_start, 1, country_headers, country_rows, "LandenTabel")
    position_start = 10
    position_headers = ["Positie", "#", "Gem. lengte", "Gem. leeftijd"]
    _write_table(ws, position_start, 11, position_headers, [[row.get(key) for key in ("position", "players", "avg_height", "avg_age")] for row in report["positions"]], "PositiesTabel")

    year_start = 18
    _write_table(ws, year_start, 1, ["Geboortejaar", "# spelers"], [[year, count] for year, count in report["birth_years"].items()], "JarenTabel")
    quarter_start = year_start
    _write_table(ws, quarter_start, 4, ["Kwartaal", "# spelers"], [[q, report["quarters"][q]] for q in ("Q1", "Q2", "Q3", "Q4")], "KwartalenTabel")

    ws["K18"] = "Rankinganalyse"
    ws["L18"] = "Waarde"
    ws["K19"] = "Ranking-lengte"
    ws["L19"] = report["correlations"]["ranking_height"] if report["correlations"]["ranking_height"] is not None else "Nog niet beschikbaar"
    ws["K20"] = "Ranking-leeftijd"
    ws["L20"] = report["correlations"]["ranking_age"] if report["correlations"]["ranking_age"] is not None else "Nog niet beschikbaar"
    for cell in ws[18][10:12]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE)
    for row in (19, 20):
        ws.cell(row, 11).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(row, 12).fill = PatternFill("solid", fgColor="FFF3E0")
        if isinstance(ws.cell(row, 12).value, float):
            ws.cell(row, 12).number_format = "0.000"
    ws.merge_cells("K21:N22")
    ws["K21"] = "Negatief betekent dat een hogere waarde samenhangt met een betere (lagere) ranking. Correlatie beschrijft geen oorzaak."
    ws["K21"].font = Font(size=9, italic=True, color="52606D")
    ws["K21"].alignment = Alignment(wrap_text=True, vertical="top")

    if report["countries"]:
        chart = _column_chart("Gem. lengte netspelers per land", "Land", "Lengte (cm)")
        chart.add_data(Reference(ws, min_col=5, min_row=country_start, max_row=country_start + len(report["countries"])), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=country_start + 1, max_row=country_start + len(report["countries"])))
        heights = [row["avg_net_height"] for row in report["countries"] if row["avg_net_height"] is not None]
        if heights:
            chart.y_axis.scaling.min = max(0, floor(min(heights) / 2) * 2 - 4)
            chart.y_axis.scaling.max = ceil(max(heights) / 2) * 2 + 2
            chart.y_axis.majorUnit = 2
        _color_points(chart, len(report["countries"]))
        ws.add_chart(chart, "H24")

    quarter_chart = _column_chart("Geboortekwartaal alle spelers", "Kwartaal", "Aantal spelers")
    quarter_chart.add_data(Reference(ws, min_col=5, min_row=quarter_start, max_row=quarter_start + 4), titles_from_data=True)
    quarter_chart.set_categories(Reference(ws, min_col=4, min_row=quarter_start + 1, max_row=quarter_start + 4))
    quarter_chart.y_axis.scaling.min = 0
    quarter_chart.y_axis.scaling.max = ceil(max(report["quarters"].values()) / 5) * 5 + 5
    _color_points(quarter_chart, 4)
    ws.add_chart(quarter_chart, "A24")

    _finish_sheet(ws, widths={"A": 11, "B": 11, "C": 11, "D": 15, "E": 20, "F": 8, "G": 8, "H": 8, "I": 8, "J": 2, "K": 18, "L": 13, "M": 15, "N": 15}, freeze="A10")
    ws.freeze_panes = None
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.print_area = "A1:N39"
    ws.oddFooter.center.text = "Lokaal gegenereerd toernooioverzicht"
    ws.sheet_properties.tabColor = NAVY


def _column_chart(title: str, x_title: str, y_title: str) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.title.overlay = False
    chart.y_axis.title = y_title
    chart.y_axis.numFmt = "0"
    chart.x_axis.title = x_title
    chart.legend = None
    chart.height = 8.0
    chart.width = 13.0
    chart.gapWidth = 55
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.showLegendKey = False
    chart.dLbls.showSerName = False
    chart.dLbls.showCatName = False
    chart.dLbls.showPercent = False
    chart.dLbls.numFmt = "0"
    chart.dLbls.dLblPos = "outEnd"
    return chart


def _color_points(chart: BarChart, count: int) -> None:
    if not chart.series:
        return
    points = []
    for index in range(count):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = PALETTE[index % len(PALETTE)]
        point.graphicalProperties.line.solidFill = PALETTE[index % len(PALETTE)]
        points.append(point)
    chart.series[0].dPt = points


def _players(ws, report: dict[str, Any]) -> None:
    frame = report["players_frame"].copy()
    headers = ["Land", "SpelerID", "Rugnr", "Achternaam", "Voornaam", "Positie", "Geb datum", "Lengte", "Leeftijd bij deelname", "Geboortejaar", "Startdatum", "Kwartaal", "Eindresultaat"]
    ranking = {team["country_code"]: team.get("final_ranking") for team in report["teams"]}
    if not frame.empty:
        frame["_output_ranking"] = frame["country_code"].map(ranking)
        frame = frame.sort_values(
            by=["_output_ranking", "country_code", "jersey_number", "id"],
            na_position="last",
            kind="stable",
        )
    rows = []
    for _, row in frame.iterrows():
        birth = row["birth_date"].date() if row["birth_date"] == row["birth_date"] else None
        player_id = f"{row['country_code']}-{int(row['id']):04d}" if "id" in row and row["id"] == row["id"] else ""
        rows.append([
            row["country_code"], player_id, row["jersey_number"], row["last_name"], row["first_name"], row["position_normalized"],
            birth, row["height_cm"], row["age"], row["birth_year"], report["tournament"]["start_date"], row["quarter"], ranking.get(row["country_code"]),
        ])
    _write_table(ws, 1, 1, headers, rows, "SpelersTabel")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{max(2, len(rows) + 1)}"
    for row in range(2, len(rows) + 2):
        ws.cell(row, 7).number_format = "dd-mm-yyyy"
        ws.cell(row, 9).number_format = "0.00"
        ws.cell(row, 11).number_format = "dd-mm-yyyy"
    _finish_sheet(ws, widths={"A": 10, "B": 16, "C": 9, "D": 25, "E": 22, "F": 11, "G": 14, "H": 10, "I": 21, "J": 14, "K": 14, "L": 11, "M": 15}, freeze="A2")


def _results(ws, report: dict[str, Any]) -> None:
    ws.merge_cells("A1:D1")
    ws["A1"] = "Resultaten en rankinganalyse"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    rows = [[r["country_code"], r["ranking"], r["avg_net_height"], r["avg_age"]] for r in report["results"]]
    _write_table(ws, 3, 1, ["Land", "Eindranking", "Gem. lengte netspelers", "Gem. leeftijd"], rows, "ResultatenTabel")
    start = 6 + len(rows)
    ws.cell(start, 1, "Analyse na eindranking")
    ws.cell(start, 1).font = Font(size=14, bold=True, color=NAVY)
    labels = [("Correlatie ranking-lengte", report["correlations"]["ranking_height"]), ("Correlatie ranking-leeftijd", report["correlations"]["ranking_age"])]
    for offset, (label, value) in enumerate(labels, 1):
        ws.cell(start + offset, 1, label)
        ws.cell(start + offset, 2, value if value is not None else "Nog niet beschikbaar")
        ws.cell(start + offset, 1).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(start + offset, 2).fill = PatternFill("solid", fgColor="FFF3E0")
        if isinstance(value, float):
            ws.cell(start + offset, 2).number_format = "0.000"
    ws.merge_cells(start_row=start + 4, start_column=1, end_row=start + 5, end_column=4)
    ws.cell(start + 4, 1, "Negatief betekent dat een hogere gemiddelde waarde samenhangt met een betere (lagere) ranking. Correlatie beschrijft samenhang, geen oorzaak.")
    ws.cell(start + 4, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[start + 4].height = 34
    _finish_sheet(ws, widths={"A": 30, "B": 20, "C": 26, "D": 18}, freeze="A4")


def _write_table(ws, start_row: int, start_col: int, headers: list[str], rows: list[list[Any]], name: str) -> None:
    for col, header in enumerate(headers, start_col):
        ws.cell(start_row, col, header)
    for row_index, values in enumerate(rows, start_row + 1):
        for col, value in enumerate(values, start_col):
            if hasattr(value, "item"):
                value = value.item()
            if value != value:
                value = None
            ws.cell(row_index, col, value)
    end_row = max(start_row + 1, start_row + len(rows))
    end_col = start_col + len(headers) - 1
    table = Table(displayName=name, ref=f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    ws.add_table(table)


def _finish_sheet(ws, widths: dict[str, float], freeze: str) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    thin = Side(style="thin", color="D9E2EC")
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal,
                    vertical="center",
                    wrap_text=cell.alignment.wrap_text,
                )
                cell.border = Border(bottom=thin)
