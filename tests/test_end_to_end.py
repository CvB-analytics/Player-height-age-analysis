from datetime import date
from io import BytesIO

import pandas as pd
import pdfplumber
from openpyxl import load_workbook

from src.parser.roster_parser import parse_bulletin
from src.reports.calculations import build_report_data
from src.reports.excel_export import create_excel_report
from src.reports.pdf_export import create_pdf_report
from src.storage.database import Database


def test_pdf_database_report_roundtrip(tmp_path, synthetic_bulletin):
    result = parse_bulletin(str(synthetic_bulletin))
    db = Database(tmp_path / "app.db")
    tournament_id = db.create_tournament(
        {
            "name": "WEVZA U20",
            "category": "U20",
            "gender": "Women",
            "location": "Terville",
            "start_date": date(2026, 8, 25),
            "source_file": synthetic_bulletin.name,
        }
    )
    db.save_players(tournament_id, result.as_records())
    teams = db.get_teams(tournament_id)
    players = db.get_players(tournament_id)
    assert len(teams) == 6
    assert len(players) == 82

    db.save_ranking(tournament_id, {team["id"]: rank for rank, team in enumerate(teams, 1)})
    report = build_report_data(db.get_tournament(tournament_id), db.get_teams(tournament_id), players)
    workbook = load_workbook(BytesIO(create_excel_report(report)))
    assert workbook.sheetnames == ["Overzicht", "Spelers", "Resultaten"]
    assert workbook["Spelers"].max_row == 83
    overview_rankings = [workbook["Overzicht"].cell(row, 2).value for row in range(11, 17)]
    result_rankings = [workbook["Resultaten"].cell(row, 2).value for row in range(4, 10)]
    assert overview_rankings == [1, 2, 3, 4, 5, 6]
    assert result_rankings == [1, 2, 3, 4, 5, 6]
    charts = workbook["Overzicht"]._charts
    assert len(charts) == 2
    assert all(chart.legend is None for chart in charts)
    assert all(chart.dLbls.showVal is True for chart in charts)
    assert all(chart.dLbls.showSerName is False for chart in charts)
    assert all(chart.dLbls.showCatName is False for chart in charts)
    assert all(chart.dLbls.numFmt == "0" for chart in charts)
    assert all(chart.layout.manualLayout.h == 0.7 for chart in charts)
    height_chart = charts[0]
    assert height_chart.x_axis.title.tx.rich.p[0].r[0].t == "Land"
    assert height_chart.y_axis.title.tx.rich.p[0].r[0].t == "Lengte (cm)"
    assert height_chart.y_axis.majorUnit == 2
    assert workbook["Overzicht"].page_setup.fitToWidth == 1
    assert workbook["Overzicht"].print_area == "'Overzicht'!$A$1:$N$39"
    pdf = create_pdf_report(report)
    assert pdf.startswith(b"%PDF")


def test_sixteen_team_exports_do_not_overlap_tables_and_charts():
    codes = ["BUL", "BEL", "ITA", "POL", "GER", "ESP", "SLO", "HUN", "TUR", "CRO", "GRE", "ROU", "NED", "FIN", "SRB", "LTU"]
    countries = [
        {
            "country_code": code,
            "players": 14,
            "avg_age": 16.1 + index / 25,
            "avg_net_height": 177.2 + (index % 8) * 1.25,
            "Q1": 3, "Q2": 4, "Q3": 3, "Q4": 4,
        }
        for index, code in enumerate(codes)
    ]
    report = {
        "tournament": {
            "name": "EK U18W 2024", "category": "U18", "gender": "Women",
            "location": "ROU", "start_date": date(2024, 7, 1), "source_file": "bulletin.pdf",
        },
        "summary": {"teams": 16, "players": 224, "avg_age": 16.51, "avg_height": 179.81},
        "countries": countries,
        "positions": [
            {"position": code, "players": 40, "avg_height": 181.0, "avg_age": 16.5}
            for code in ("DIA", "LIB", "MB", "PL", "SV")
        ],
        "birth_years": {
            1993: 3, 1994: 4, 1995: 4, 1996: 2, 1997: 10, 1998: 11,
            1999: 12, 2000: 13, 2001: 14, 2002: 15, 2003: 16, 2004: 17,
            2005: 18, 2006: 19, 2007: 20, 2008: 21, 2009: 12, 2010: 3,
        },
        "quarters": {"Q1": 75, "Q2": 63, "Q3": 43, "Q4": 43},
        "correlations": {"ranking_height": -0.501, "ranking_age": 0.290},
        "results": [
            {"country_code": code, "ranking": index + 1, "avg_net_height": countries[index]["avg_net_height"], "avg_age": countries[index]["avg_age"]}
            for index, code in enumerate(codes)
        ],
        "teams": [{"country_code": code, "final_ranking": index + 1} for index, code in enumerate(codes)],
        "players_frame": pd.DataFrame(),
    }

    workbook = load_workbook(BytesIO(create_excel_report(report)))
    overview = workbook["Overzicht"]
    country_end = overview.tables["LandenTabel"].ref.split(":")[1]
    year_start = overview.tables["JarenTabel"].ref.split(":")[0]
    year_end = overview.tables["JarenTabel"].ref.split(":")[1]
    assert int("".join(filter(str.isdigit, year_start))) > int("".join(filter(str.isdigit, country_end)))
    assert all(
        chart.anchor._from.row + 1 > int("".join(filter(str.isdigit, year_end)))
        for chart in overview._charts
    )
    assert overview.page_setup.fitToWidth == 1
    assert overview.page_setup.fitToHeight == 2

    pdf_bytes = create_pdf_report(report)
    with pdfplumber.open(BytesIO(pdf_bytes)) as document:
        text = document.pages[0].extract_text() or ""
    assert all(code in text for code in codes)
    assert "Aantal spelers per geboortejaar" in text
    assert "1993" in text
    assert "2010" in text
